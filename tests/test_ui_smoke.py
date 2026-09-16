"""Headless construction tests for the primary Qt surfaces."""

from __future__ import annotations

import dataclasses
import threading
import time
from collections.abc import Callable
from datetime import UTC, date, datetime, timedelta
from pathlib import Path

import pytest
from PySide6.QtCore import QDate
from PySide6.QtGui import QCursor, QFont, QGuiApplication, QPalette
from PySide6.QtWidgets import (
    QApplication,
    QButtonGroup,
    QCheckBox,
    QComboBox,
    QDialog,
    QFrame,
    QLabel,
    QMessageBox,
    QPushButton,
    QSystemTrayIcon,
)

from organizador import __version__, duplicates, updater
from organizador.classifier import guess_filing
from organizador.config import AppConfig
from organizador.controller import AppController, _UndoJob
from organizador.db import Database
from organizador.filer import FilingService
from organizador.models import ExistingDownload, FiledDocument, FindingReason, Subject
from organizador.paths import IncompleteMoveError
from organizador.reconcile import findings, visible_findings
from organizador.reconcile import scan as scan_reconciliation
from organizador.recovery import (
    PRE_RESTORE_MARKER,
    RESTORE_REQUEST_NAME,
    USER_MARKER,
    BundleInfo,
)
from organizador.ui.dialogs import (
    BackupListDialog,
    MoveDocumentDialog,
    OnboardingDialog,
    SubjectDialog,
    SubjectFilesDialog,
    TaskDialog,
)
from organizador.ui.main_window import MainWindow
from organizador.ui.pages import SettingsPage, SettingsPayload
from organizador.ui.prompt import FilingPrompt
from organizador.ui.theme import apply_theme, get_theme
from organizador.ui.tray import TrayIcon
from organizador.ui.widgets import EmptyState


def test_main_window_builds_and_refreshes_all_pages(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    apply_theme(qt_app, get_theme("escuro"))
    window = MainWindow(database, app_config)

    window.refresh_all(watching=True, paused=False)
    window.show_page("disciplinas")
    qt_app.processEvents()

    assert window.stack.currentWidget() is window.subjects_page
    assert window.status_label.text() == "A vigiar Downloads"
    assert (
        qt_app.palette().color(QPalette.ColorRole.Window).name()
        == get_theme("escuro").canvas.casefold()
    )
    visible_copy = [item.text() for item in window.subjects_page.findChildren(QLabel)]
    assert any(subject.name in text for text in visible_copy)
    settings_copy = [item.text() for item in window.settings_page.findChildren(QLabel)]
    assert any(f"Organizador v{__version__}" in text for text in settings_copy)
    import_requests: list[bool] = []
    window.inbox_page.import_existing_requested.connect(lambda: import_requests.append(True))
    window.inbox_page.import_button.click()
    assert import_requests == [True]
    window.allow_close = True
    window.close()


def test_settings_preserve_an_exact_minimum_file_size(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
) -> None:
    del qt_app
    app_config.minimum_file_size = 1537
    window = MainWindow(database, app_config)
    payloads: list[SettingsPayload] = []
    window.settings_page.save_requested.connect(payloads.append)

    save = next(
        control
        for control in window.settings_page.findChildren(QPushButton)
        if control.text() == "Guardar definições"
    )
    save.click()

    assert payloads[0]["minimum_file_size"] == 1537
    window.allow_close = True
    window.close()


def test_settings_ocr_checkbox_round_trips_through_payload(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
) -> None:
    del qt_app
    app_config.ocr_enabled = False
    window = MainWindow(database, app_config)
    payloads: list[SettingsPayload] = []
    window.settings_page.save_requested.connect(payloads.append)

    ocr_box = next(
        control
        for control in window.settings_page.findChildren(QCheckBox)
        if control.text() == "Reconhecer texto em PDFs digitalizados (OCR)"
    )
    assert not ocr_box.isChecked()
    ocr_box.setChecked(True)
    save = next(
        control
        for control in window.settings_page.findChildren(QPushButton)
        if control.text() == "Guardar definições"
    )
    save.click()

    assert payloads[0]["ocr_enabled"] is True
    window.allow_close = True
    window.close()


def test_settings_page_gates_the_popup_timeout(
    qt_app: QApplication,
    app_config: AppConfig,
) -> None:
    page = SettingsPage(app_config)
    try:
        assert not page.auto_close_check.isChecked()
        assert not page.timeout_spin.isEnabled()

        page.auto_close_check.setChecked(True)
        assert page.timeout_spin.isEnabled()

        app_config.prompt_timeout_seconds = 90
        app_config.prompt_timeout_enabled = True
        page.load_config(app_config)
        assert page.auto_close_check.isChecked()
        assert page.timeout_spin.isEnabled()
        assert page.timeout_spin.value() == 90

        payloads: list[SettingsPayload] = []
        page.save_requested.connect(payloads.append)
        page._save()
        assert payloads and payloads[0]["prompt_timeout_enabled"] is True
        assert payloads[0]["prompt_timeout_seconds"] == 90
    finally:
        page.deleteLater()
        qt_app.processEvents()


def _bundle_info(
    path: str, kind: str, when: datetime | None = None, size: int = 2048
) -> BundleInfo:
    return BundleInfo(
        path=Path(path),
        created_at=when or datetime(2026, 9, 13, 12, tzinfo=UTC),
        database_user_version=6,
        settings_present=True,
        kind=kind,
        size_bytes=size,
    )


def test_settings_backup_panel_summarizes_snapshots(
    qt_app: QApplication,
    app_config: AppConfig,
) -> None:
    del qt_app
    page = SettingsPage(app_config)
    try:
        page.set_backup_inventory(())
        assert "Ainda não há cópias" in page.backup_status.text()

        page.set_backup_inventory((_bundle_info("C:/b/user-a", USER_MARKER),))
        assert "Última cópia" in page.backup_status.text()

        page.set_backup_inventory(
            (
                _bundle_info("C:/b/user-a", USER_MARKER),
                _bundle_info("C:/b/user-b", USER_MARKER),
            )
        )
        assert "2 cópias" in page.backup_status.text()

        page.set_backup_busy(True)
        assert not page.backup_create_button.isEnabled()
        assert not page.backup_export_button.isEnabled()
        assert not page.backup_manage_button.isEnabled()
        assert "A processar" in page.backup_status.text()

        page.set_backup_busy(False)
        assert page.backup_create_button.isEnabled()

        requested: list[str] = []
        page.backup_created_requested.connect(lambda: requested.append("create"))
        page.backup_export_requested.connect(lambda: requested.append("export"))
        page.backup_manager_requested.connect(lambda: requested.append("manage"))
        page.backup_open_folder_requested.connect(lambda: requested.append("folder"))
        page.backup_create_button.click()
        page.backup_export_button.click()
        page.backup_manage_button.click()
        page.backup_folder_button.click()

        assert requested == ["create", "export", "manage", "folder"]
    finally:
        page.deleteLater()


def test_backup_list_dialog_actions_respect_bundle_kinds(qt_app: QApplication) -> None:
    del qt_app
    user = _bundle_info("C:/b/user-a", USER_MARKER, datetime(2026, 9, 13, 12, tzinfo=UTC))
    migration = _bundle_info("C:/b/migration-a", "migration", datetime(2026, 9, 12, 10, tzinfo=UTC))
    pre_restore = _bundle_info("C:/b/pre-restore-a", PRE_RESTORE_MARKER)
    dialog = BackupListDialog((user, migration, pre_restore))
    try:
        buttons = dialog.findChildren(QPushButton)
        restore_buttons = [control for control in buttons if control.text() == "Restaurar"]
        export_buttons = [control for control in buttons if control.text() == "Exportar"]
        delete_buttons = [control for control in buttons if control.text() == "Remover"]
        assert len(restore_buttons) == 3
        assert len(export_buttons) == 3
        assert len(delete_buttons) == 1

        restored: list[BundleInfo] = []
        dialog.restore_requested.connect(restored.append)
        for control in restore_buttons:
            control.click()

        assert restored == [user, migration, pre_restore]

        removed: list[BundleInfo] = []
        dialog.delete_requested.connect(removed.append)
        delete_buttons[0].click()

        assert removed == [user]

        dialog.set_bundles(())
        assert any(
            control.text() == "Importar .zip…" for control in dialog.findChildren(QPushButton)
        )
    finally:
        dialog.deleteLater()


def test_controller_creates_backups_and_stages_restores(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    page = controller.main_window.settings_page
    try:
        controller._pending_transfers = 1
        controller.create_user_backup()
        assert "Espera" in page.backup_status.text()
        controller._pending_transfers = 0

        controller.create_user_backup()
        _pump_until(qt_app, lambda: page.backup_create_button.isEnabled())

        bundles = controller._list_bundles()
        user_bundle = next(info for info in bundles if info.kind == USER_MARKER)
        assert "criada" in page.backup_status.text()

        monkeypatch.setattr(
            QMessageBox, "question", lambda *args, **kwargs: QMessageBox.StandardButton.Yes
        )
        relaunches: list[bool] = []
        monkeypatch.setattr(controller, "_relaunch_for_restore", lambda: relaunches.append(True))

        controller._request_restore(user_bundle)

        assert relaunches == [True]
        assert (app_config.data_dir / RESTORE_REQUEST_NAME).is_file()

        dialog = BackupListDialog((user_bundle,), controller.main_window)
        try:
            controller._delete_backup(user_bundle, dialog)
        finally:
            dialog.deleteLater()

        assert not user_bundle.path.exists()
        assert controller._list_bundles() == ()
    finally:
        (app_config.data_dir / RESTORE_REQUEST_NAME).unlink(missing_ok=True)
        _close_controller(qt_app, controller)


def test_settings_quiet_checkbox_round_trips_through_payload(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
) -> None:
    del qt_app
    app_config.quiet_intake = True
    window = MainWindow(database, app_config)
    payloads: list[SettingsPayload] = []
    window.settings_page.save_requested.connect(payloads.append)

    quiet_box = next(
        control
        for control in window.settings_page.findChildren(QCheckBox)
        if control.text() == "Silenciar notificações de arquivo"
    )
    assert quiet_box.isChecked()
    quiet_box.setChecked(False)
    save = next(
        control
        for control in window.settings_page.findChildren(QPushButton)
        if control.text() == "Guardar definições"
    )
    save.click()

    assert payloads[0]["quiet_intake"] is False
    window.allow_close = True
    window.close()


def test_recovery_row_offers_only_a_safe_folder_action(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    del subject
    missing_path = app_config.inbox_dir / "em-recuperacao.pdf"
    item = database.add_inbox_item(
        missing_path,
        app_config.downloads_dir / missing_path.name,
        missing_path.name,
        200,
    )
    assert database.mark_inbox_recovery_required(item.id, "Recuperação necessária")
    window = MainWindow(database, app_config)
    opened: list[object] = []
    window.inbox_page.open_path.connect(opened.append)

    window.inbox_page.refresh()
    buttons = window.inbox_page.findChildren(QPushButton)
    button_texts = {control.text() for control in buttons}

    assert window.inbox_page.summary_label.text() == "1 ficheiro precisa de recuperação manual"
    assert "Abrir Universidade" in button_texts
    assert "Organizar" not in button_texts
    assert "Não é da universidade" not in button_texts
    next(control for control in buttons if control.text() == "Abrir Universidade").click()
    assert opened == [app_config.university_root]
    window.allow_close = True
    window.close()


def test_unresolved_history_path_is_visible_in_the_inbox(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    untracked = app_config.university_root / subject.folder_name / "Slides" / "sem-registo.pdf"
    untracked.write_bytes(b"untracked but untouched")
    window = MainWindow(database, app_config)
    adopted: list[object] = []
    reviewed: list[object] = []
    window.inbox_page.adopt_requested.connect(adopted.append)
    window.inbox_page.dismiss_finding_requested.connect(reviewed.append)
    window.inbox_page.set_reconciliation_report(scan_reconciliation(app_config, database))

    window.inbox_page.refresh()
    visible_copy = [item.text() for item in window.inbox_page.findChildren(QLabel)]
    buttons = window.inbox_page.findChildren(QPushButton)

    assert "1 ocorrência do histórico precisa de revisão" in window.inbox_page.summary_label.text()
    assert any("Encontrado numa disciplina sem registo" in text for text in visible_copy)
    assert any(str(untracked) == text for text in visible_copy)
    next(control for control in buttons if control.text() == "Adotar").click()
    next(control for control in buttons if control.text() == "Marcar revisto").click()
    assert len(adopted) == 1
    assert len(reviewed) == 1
    assert untracked.read_bytes() == b"untracked but untouched"
    window.allow_close = True
    window.close()


def test_two_reconciliation_reasons_at_one_path_render_as_two_findings(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    filer: FilingService,
    subject: Subject,
) -> None:
    source = app_config.downloads_dir / "same-path.txt"
    source.write_bytes(b"one path with two reasons")
    item = filer.ingest(source)
    assert item is not None
    document = filer.file_document(item.id, subject.id, "Outros", source.name)
    document.current_path.unlink()
    window = MainWindow(database, app_config)
    window.inbox_page.set_reconciliation_report(scan_reconciliation(app_config, database))

    window.inbox_page.refresh()
    visible_copy = [item.text() for item in window.inbox_page.findChildren(QLabel)]
    button_copy = [item.text() for item in window.inbox_page.findChildren(QPushButton)]

    assert (
        "2 ocorrências do histórico precisam de revisão" in window.inbox_page.summary_label.text()
    )
    assert any("Documento registado" in text for text in visible_copy)
    assert any("não pode ser desfeita" in text for text in visible_copy)
    assert visible_copy.count(str(document.current_path)) == 2
    assert "Remover registo" in button_copy
    assert button_copy.count("Marcar revisto") == 2
    window.allow_close = True
    window.close()


def test_controller_adopts_and_unregisters_without_moving_the_file(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    path = app_config.university_root / subject.folder_name / "Outros" / "legacy.txt"
    contents = b"catalog this file in place"
    path.write_bytes(contents)
    controller = AppController(app_config)
    report = scan_reconciliation(app_config, controller.database)
    finding = next(
        item for item in findings(report) if item.reason is FindingReason.UNTRACKED_SUBJECT_FILE
    )
    monkeypatch.setattr(
        QMessageBox,
        "question",
        lambda *_args, **_kwargs: QMessageBox.StandardButton.Yes,
    )

    controller._adopt_untracked_file(finding)
    adopted = controller.database.list_adopted_files()

    assert len(adopted) == 1
    assert path.read_bytes() == contents
    controller._unregister_adopted_file(adopted[0].id)
    report = scan_reconciliation(app_config, controller.database)
    assert controller.database.list_adopted_files() == []
    assert visible_findings(controller.database, report) == ()
    assert path.read_bytes() == contents
    qt_app.processEvents()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_home_activity_panel_reports_safety_history(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    inbox_path = app_config.inbox_dir / "MAT101_historico.pdf"
    inbox_path.write_bytes(b"activity history" * 20)
    item = database.add_inbox_item(
        inbox_path,
        app_config.downloads_dir / inbox_path.name,
        inbox_path.name,
        inbox_path.stat().st_size,
    )
    existing = app_config.university_root / subject.folder_name / "Slides" / "MAT101_historico.pdf"
    existing.write_bytes(b"already filed here")
    filer = FilingService(app_config, database)
    filer.file_document(item.id, subject.id, "Slides", "MAT101_historico.pdf")
    window = MainWindow(database, app_config)

    window.refresh_all(watching=True, paused=False)

    text = window.home_page.activity_label.text()
    assert "1 ficheiro organizado" in text
    assert "1 colisão de nomes resolvida" in text
    window.allow_close = True
    window.close()


def test_home_activity_panel_starts_with_a_calm_empty_state(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
) -> None:
    window = MainWindow(database, app_config)

    window.refresh_all(watching=True, paused=False)

    assert "ainda não tem histórico" in window.home_page.activity_label.text()
    window.allow_close = True
    window.close()


def test_stale_watcher_generation_cannot_ingest_after_restart(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
) -> None:
    del qt_app, database
    controller = AppController(app_config)
    controller._watcher_generation = 2
    candidate = app_config.downloads_dir / "stale-callback.pdf"
    contents = b"stale watcher callback must be ignored"
    candidate.write_bytes(contents)

    controller._ingest_download(1, candidate)

    assert controller.database.count_inbox_items() == 0
    assert candidate.read_bytes() == contents
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_filing_prompt_selects_classifier_suggestion(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    path = app_config.inbox_dir / "MAT101_ficha.pdf"
    path.write_bytes(b"content")
    item = database.add_inbox_item(path, app_config.downloads_dir / path.name, path.name, 7)
    guess = guess_filing(item.original_name, [subject])
    prompt = FilingPrompt(timeout_seconds=30)

    prompt.show_item(item, [subject], guess)
    qt_app.processEvents()

    assert prompt.current_item_id == item.id
    assert prompt.selected_subject_id == subject.id
    assert prompt.confirm_button.isEnabled()
    assert prompt.type_buttons["Exercícios"].isChecked()
    screen = QGuiApplication.screenAt(QCursor.pos()) or prompt.screen()
    target = prompt.animation.endValue()
    assert target.x() == screen.availableGeometry().left() + 18

    prompt.show_item(item, [subject], guess)
    qt_app.processEvents()
    assert len(prompt.findChildren(QButtonGroup)) == 2
    prompt.timer.stop()
    prompt.hide()


def test_filing_prompt_regenerates_name_on_subject_and_type_change(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    other = database.add_subject(
        "Biologia",
        "BIO101",
        "#123456",
        (),
        "BIO101 - Biologia",
    )
    subjects = [subject, other]
    path = app_config.inbox_dir / "BIO101_ficha.pdf"
    path.write_bytes(b"content")
    item = database.add_inbox_item(path, app_config.downloads_dir / path.name, path.name, 7)
    guess = guess_filing(item.original_name, subjects)
    prompt = FilingPrompt(timeout_seconds=30)
    try:
        prompt.show_item(item, subjects, guess, name_template="{codigo}_{tipo}_{nome_original}")
        qt_app.processEvents()
        initial = prompt.name_edit.text()
        assert other.code in initial

        prompt._choose_subject(subject.id, prompt.subject_group.button(subject.id))
        qt_app.processEvents()
        regenerated = prompt.name_edit.text()
        assert regenerated != initial
        assert regenerated.startswith(f"{subject.code}_")
        assert initial.startswith(f"{other.code}_")

        prompt.type_buttons["Slides"].click()
        qt_app.processEvents()
        assert prompt.type_buttons["Slides"].isChecked()
        assert "Slides" in prompt.name_edit.text()

        prompt.name_edit.setText("o meu nome.pdf")
        prompt.type_buttons["Trabalhos"].click()
        qt_app.processEvents()
        assert prompt.name_edit.text() == "o meu nome.pdf"
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_filing_prompt_offers_replacement_only_on_the_matching_folder(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    folder = app_config.university_root / subject.folder_name / "Slides"
    folder.mkdir(parents=True, exist_ok=True)
    old_path = folder / "Aula.pdf"
    old_path.write_bytes(b"versao antiga do documento")
    candidate = ExistingDownload.capture(old_path)
    assert candidate is not None
    old = database.adopt_subject_file(candidate, subject.id, "Slides")

    inbox_path = app_config.inbox_dir / "Aula.pdf"
    inbox_path.write_bytes(b"versao antiga do documento")
    item = database.add_inbox_item(
        inbox_path,
        app_config.downloads_dir / inbox_path.name,
        inbox_path.name,
        inbox_path.stat().st_size,
    )
    duplicate = duplicates.find_duplicate(database, item)
    assert duplicate is not None
    prompt = FilingPrompt(timeout_seconds=30)
    try:
        guess = guess_filing(item.original_name, [subject])
        prompt.show_item(item, [subject], guess, duplicate=duplicate)
        qt_app.processEvents()

        assert prompt.duplicate_banner.isVisible()
        prompt._choose_subject(subject.id, prompt.subject_group.button(subject.id))
        prompt.type_buttons["Slides"].click()
        qt_app.processEvents()
        assert prompt.replace_check.isEnabled()

        revealed: list[object] = []
        prompt.reveal_requested.connect(revealed.append)
        prompt._reveal_duplicate()
        assert revealed == [old.current_path]

        prompt.type_buttons["Testes"].click()
        qt_app.processEvents()
        assert not prompt.replace_check.isEnabled()
        assert not prompt.replace_check.isChecked()

        prompt.type_buttons["Slides"].click()
        qt_app.processEvents()
        prompt.replace_check.setChecked(True)
        emitted: list[tuple[object, ...]] = []
        prompt.filing_requested.connect(lambda *args: emitted.append(args))
        prompt._confirm()

        assert emitted and emitted[0][6] == old.id
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_prompt_flow_detects_an_existing_copy(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    folder = app_config.university_root / subject.folder_name / "Slides"
    folder.mkdir(parents=True, exist_ok=True)
    old_path = folder / "Aula.pdf"
    old_path.write_bytes(b"conteudo repetido para o teste")
    candidate = ExistingDownload.capture(old_path)
    assert candidate is not None
    database.adopt_subject_file(candidate, subject.id, "Slides")

    inbox_path = app_config.inbox_dir / "Aula.pdf"
    inbox_path.write_bytes(b"conteudo repetido para o teste")
    item = database.add_inbox_item(
        inbox_path,
        app_config.downloads_dir / inbox_path.name,
        inbox_path.name,
        inbox_path.stat().st_size,
    )

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller.prompt_queue.append(item.id)
        controller._show_next_prompt()
        qt_app.processEvents()

        assert controller.prompt.current_item_id == item.id
        assert controller.prompt._duplicate is not None
        assert controller.prompt._duplicate.current_path == old_path
        assert controller.prompt.duplicate_banner.isVisible()
    finally:
        controller.prompt.timer.stop()
        controller.prompt.hide()
        _close_controller(qt_app, controller)


def test_filing_prompt_elides_long_subject_names(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    long_name = "Armazenamento para Bibliotecas Digitais"
    long_subject = database.add_subject(long_name, "ABD", "#087A74", (), "ABD - Armazenamento")
    inbox_path = app_config.inbox_dir / "aula.pdf"
    inbox_path.write_bytes(b"conteudo para o teste de elisao")
    item = database.add_inbox_item(
        inbox_path,
        app_config.downloads_dir / inbox_path.name,
        inbox_path.name,
        inbox_path.stat().st_size,
    )
    prompt = FilingPrompt(timeout_seconds=30)
    try:
        subjects = [subject, long_subject]
        guess = guess_filing(item.original_name, subjects)
        prompt.show_item(item, subjects, guess)
        qt_app.processEvents()
        prompt.adjustSize()
        qt_app.processEvents()

        long_button = prompt.subject_group.button(long_subject.id)
        short_button = prompt.subject_group.button(subject.id)
        assert long_button is not None and short_button is not None
        assert long_button.full_text() == f"2  {long_name}"
        visible = long_button.text().rstrip("…")
        assert long_button.text().endswith("…")
        assert len(long_button.text()) < len(f"2  {long_name}")
        assert visible.startswith("2  Armazena")
        assert f"2  {long_name}".startswith(visible)
        assert long_button.accessibleName() == f"2  {long_name}"
        assert long_name in long_button.toolTip()
        assert short_button.text() == f"1  {subject.name}"

        prompt._choose_subject(long_subject.id, long_button)
        assert prompt.selected_subject_id == long_subject.id
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_filing_prompt_waits_for_a_decision_by_default(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    inbox_path = app_config.inbox_dir / "aula.pdf"
    inbox_path.write_bytes(b"conteudo para o teste do temporizador")
    item = database.add_inbox_item(
        inbox_path,
        app_config.downloads_dir / inbox_path.name,
        inbox_path.name,
        inbox_path.stat().st_size,
    )
    prompt = FilingPrompt(timeout_seconds=30)
    returned: list[int] = []
    later: list[int] = []
    prompt.return_requested.connect(returned.append)
    prompt.later_requested.connect(later.append)
    try:
        guess = guess_filing(item.original_name, [subject])
        prompt.show_item(item, [subject], guess)
        qt_app.processEvents()

        assert not prompt.timer.isActive()
        assert not prompt.countdown_label.isVisible()
        tooltip = "Não é da universidade: devolve o ficheiro"
        assert prompt.close_button.toolTip() == tooltip
        assert prompt.close_button.accessibleName() == tooltip

        prompt.close_button.click()
        qt_app.processEvents()

        assert returned == [item.id]
        assert later == []
        assert prompt.current_item_id is None

        prompt.show_item(item, [subject], guess)
        qt_app.processEvents()
        later_button = next(
            control
            for control in prompt.findChildren(QPushButton)
            if control.text() == "Mais tarde"
        )
        later_button.click()
        qt_app.processEvents()

        assert later == [item.id]
        assert returned == [item.id]
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_prompt_close_returns_the_file_to_its_origin(
    qt_app: QApplication,
    app_config: AppConfig,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    del subject
    origin = tmp_path / "Ambiente de Trabalho"
    origin.mkdir()
    source = origin / "apontamento.pdf"
    source.write_bytes(b"external note content " * 8)
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller.organize_external_paths([str(source)])
        _pump_until(qt_app, lambda: controller.prompt.current_item_id is not None)

        controller.prompt.close_button.click()
        _pump_until(qt_app, lambda: any(args[0] == "Ficheiro devolvido" for args, _ in notices))

        assert source.read_bytes() == b"external note content " * 8
        assert controller.prompt.current_item_id is None
        assert controller.database.count_inbox_items() == 0
    finally:
        controller.prompt.timer.stop()
        controller.prompt.hide()
        _close_controller(qt_app, controller)


def test_filing_prompt_counts_down_when_auto_close_is_enabled(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    inbox_path = app_config.inbox_dir / "aula.pdf"
    inbox_path.write_bytes(b"conteudo para o teste do temporizador")
    item = database.add_inbox_item(
        inbox_path,
        app_config.downloads_dir / inbox_path.name,
        inbox_path.name,
        inbox_path.stat().st_size,
    )
    prompt = FilingPrompt(timeout_seconds=30, auto_close=True)
    later: list[int] = []
    prompt.later_requested.connect(later.append)
    try:
        guess = guess_filing(item.original_name, [subject])
        prompt.show_item(item, [subject], guess)
        qt_app.processEvents()

        assert prompt.timer.isActive()
        assert prompt.countdown_label.isVisible()
        assert "30s" in prompt.countdown_label.text()

        prompt.remaining = 1
        prompt._tick()
        qt_app.processEvents()

        assert later == [item.id]
        assert not prompt.timer.isActive()
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_subject_colour_button_keeps_readable_text(qt_app: QApplication) -> None:
    dialog = SubjectDialog()

    assert "color: #08111d" in dialog.color_button.styleSheet()
    dialog.color = "#08111D"
    dialog._update_color_button()
    assert "color: #ffffff" in dialog.color_button.styleSheet().casefold()
    dialog.close()


def test_empty_state_reserves_height_for_wrapped_body(qt_app: QApplication) -> None:
    apply_theme(qt_app, get_theme("escuro"))
    empty = EmptyState(
        "Tudo no lugar",
        "Quando terminares um download elegível, ele aparece aqui e num pequeno popup.",
    )
    empty.resize(900, 270)
    empty.show()
    qt_app.processEvents()

    detail = next(
        child for child in empty.findChildren(QLabel) if child.objectName() == "PageSubtitle"
    )
    assert detail.width() == 520
    assert detail.height() >= detail.heightForWidth(detail.width())
    empty.close()


def test_onboarding_removes_subject_when_settings_save_fails(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    filer: FilingService,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    app_config.initialized = False
    old_root = app_config.university_root
    dialog = OnboardingDialog(app_config, database, filer)
    dialog.root_edit.setText(str(old_root / "Nova"))
    dialog.name_edit.setText("Álgebra Linear")

    def fail_save(_config: AppConfig) -> None:
        raise OSError("disco indisponível")

    monkeypatch.setattr(AppConfig, "save", fail_save)
    dialog._finish()
    qt_app.processEvents()

    assert database.count_subjects() == 0
    assert app_config.university_root == old_root
    assert not app_config.initialized
    assert "disco indisponível" in dialog.error_label.text()
    dialog.close()


def test_startup_registration_is_reverted_when_settings_save_fails(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller = AppController(app_config)
    calls: list[bool] = []
    monkeypatch.setattr("organizador.controller.set_launch_at_login", calls.append)

    def fail_save(_config: AppConfig) -> None:
        raise OSError("sem espaço")

    monkeypatch.setattr(AppConfig, "save", fail_save)
    payload: SettingsPayload = {
        "university_root": app_config.university_root,
        "downloads_dir": app_config.downloads_dir,
        "extensions": ".pdf",
        "filename_template": "{nome_original}",
        "minimum_file_size": 1024,
        "prompt_timeout_seconds": 45,
        "prompt_timeout_enabled": True,
        "reminder_lead_days": 2,
        "theme": "escuro",
        "language": "pt",
        "check_updates_on_launch": True,
        "ocr_enabled": True,
        "quiet_intake": True,
        "watch_enabled": True,
        "launch_at_login": True,
    }

    controller._save_settings(payload)
    qt_app.processEvents()

    assert calls == [True, False]
    assert not app_config.launch_at_login
    assert not app_config.quiet_intake
    assert "sem espaço" in controller.main_window.settings_page.status_label.text()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_startup_reconciles_before_watcher_and_indexer(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database, subject
    orphan = app_config.inbox_dir / "MAT101_arranque.pdf"
    orphan.write_bytes(b"startup orphan university document")
    controller = AppController(app_config)
    order: list[str] = []

    def observe_watcher_start() -> None:
        recovered = controller.database.find_active_inbox_by_path(orphan)
        assert recovered is not None
        assert recovered.status == "pending"
        order.append("watcher")

    monkeypatch.setattr(controller, "_restart_watcher", observe_watcher_start)
    monkeypatch.setattr(controller.indexer, "submit_pending", lambda: order.append("indexer"))

    controller.start(smoke_test=True)
    qt_app.processEvents()

    recovered = controller.database.find_active_inbox_by_path(orphan)
    assert recovered is not None
    assert recovered.suggested_subject_id is not None
    assert order == ["watcher", "indexer"]
    controller.reminder_timer.stop()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_smoke_start_skips_pending_update_notifications(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database, subject
    controller = AppController(app_config)
    shown: list[bool] = []
    monkeypatch.setattr(controller, "_restart_watcher", lambda: None)
    monkeypatch.setattr(controller, "_show_pending_update_result", lambda: shown.append(True))

    controller.start(smoke_test=True)
    qt_app.processEvents()

    assert shown == []
    controller.reminder_timer.stop()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_incomplete_return_is_ignored_by_the_live_watcher(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database, subject
    controller = AppController(app_config)
    source = app_config.downloads_dir / "devolucao-parcial.pdf"
    source.write_bytes(b"return source" * 20)
    item = controller.filer.ingest(source)
    assert item is not None

    class WatcherStub:
        running = True
        paused = False

        def __init__(self) -> None:
            self.ignored: list[tuple[object, float]] = []

        def set_paused(self, paused: bool) -> None:
            self.paused = paused

        def ignore_self_move(self, path: object, *, seconds: float = 30.0) -> None:
            self.ignored.append((path, seconds))

    watcher = WatcherStub()
    controller.watcher = watcher  # type: ignore[assignment]

    def leave_partial(_source: object, target: object, **_kwargs: object) -> object:
        destination = Path(str(target))
        destination.write_bytes(b"partial return copy")
        raise IncompleteMoveError(destination)

    monkeypatch.setattr("organizador.filer.move_without_overwrite", leave_partial)

    controller._return_item(item.id)

    _pump_until(qt_app, lambda: watcher.ignored != [])

    pending = controller.database.list_pending_returns()
    assert len(pending) == 1
    assert watcher.ignored == [(pending[0].destination_path, float("inf"))]
    assert pending[0].destination_path.read_bytes() == b"partial return copy"
    assert item.path.exists()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_confirmed_existing_download_import_is_capped_and_uses_normal_inbox_flow(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database, subject
    app_config.watch_enabled = False
    for index in range(30):
        path = app_config.downloads_dir / f"existing_{index:02}.pdf"
        path.write_bytes(b"existing university material")
    confirmation_copy: list[str] = []
    answers = [QMessageBox.StandardButton.Cancel, QMessageBox.StandardButton.Yes]

    def confirm(*args: object, **_kwargs: object) -> QMessageBox.StandardButton:
        confirmation_copy.append(str(args[2]))
        return answers.pop(0)

    monkeypatch.setattr(QMessageBox, "question", confirm)
    monkeypatch.setattr("organizador.watcher.wait_until_stable", lambda *_args, **_kwargs: True)
    controller = AppController(app_config)
    controller._restart_watcher()
    controller._import_existing_downloads()
    assert controller.database.count_inbox_items() == 0
    assert len(list(app_config.downloads_dir.glob("*.pdf"))) == 30

    controller._import_existing_downloads()

    deadline = time.monotonic() + 20.0
    while controller._manual_import_active and time.monotonic() < deadline:
        qt_app.processEvents()
        time.sleep(0.01)
    qt_app.processEvents()

    assert not controller._manual_import_active
    assert controller.database.count_inbox_items() == 25
    assert len(list(app_config.downloads_dir.glob("*.pdf"))) == 5
    assert len(confirmation_copy) == 2
    assert "30 ficheiros elegíveis" in confirmation_copy[1]
    assert "no máximo 25" in confirmation_copy[1]
    assert "25 importados" in controller.main_window.inbox_page.import_status_label.text()
    assert controller.prompt.current_item_id is not None

    if controller.watcher is not None:
        controller.watcher.stop()
    controller.prompt.timer.stop()
    controller.prompt.hide()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_advance_reminders_fire_once_per_day_and_survive_restart(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    controller = AppController(app_config)
    sent: list[tuple[str, str]] = []
    monkeypatch.setattr(controller.tray, "notify", lambda title, body: sent.append((title, body)))
    soon_id = controller.database.add_task(
        "Entrega breve", subject.id, date.today() + timedelta(days=2)
    ).id
    controller.database.add_task("Longínqua", subject.id, date.today() + timedelta(days=30))

    controller._check_deadlines()
    controller._check_deadlines()

    assert len(sent) == 1
    assert "Entrega breve" in sent[0][1]

    restarted = AppController(app_config)
    restarted_sent: list[tuple[str, str]] = []
    monkeypatch.setattr(
        restarted.tray, "notify", lambda title, body: restarted_sent.append((title, body))
    )
    restarted._check_deadlines()
    assert restarted_sent == []

    controller.database.update_task(soon_id, "Entrega breve", subject.id, date.today())
    controller._check_deadlines()

    assert len(sent) == 2
    assert "vence hoje" in sent[1][1]

    controller.reminder_timer.stop()
    restarted.reminder_timer.stop()
    for active in (controller, restarted):
        active.indexer.shutdown()
        active.tray.hide()
        active.main_window.allow_close = True
        active.main_window.close()


def test_per_task_reminder_lead_overrides_the_global_default(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    controller = AppController(app_config)
    sent: list[str] = []
    monkeypatch.setattr(controller.tray, "notify", lambda title, body: sent.append(body))
    controller.database.add_task(
        "Projeto final", subject.id, date.today() + timedelta(days=5), reminder_lead_days=7
    )
    controller.database.add_task("Normal", subject.id, date.today() + timedelta(days=5))

    controller._check_deadlines()

    assert len(sent) == 1
    assert "Projeto final" in sent[0]

    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


class _StubBulkDialog:
    values: tuple[int, str, bool, date | None] = (0, "Slides", False, None)

    def __init__(self, *_args: object, **_kwargs: object) -> None:
        pass

    def exec(self) -> QDialog.DialogCode:
        return QDialog.DialogCode.Accepted


def test_bulk_filing_files_selection_and_keeps_failures_pending(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    app_config.watch_enabled = False
    controller = AppController(app_config)
    ids: list[int] = []
    for name in ("lote_um.pdf", "lote_dois.pdf", "lote_tres.pdf"):
        path = app_config.inbox_dir / name
        path.write_bytes(b"bulk content" * 10)
        item = controller.database.add_inbox_item(
            path, app_config.downloads_dir / name, name, path.stat().st_size
        )
        ids.append(item.id)
    (app_config.inbox_dir / "lote_tres.pdf").unlink()
    _StubBulkDialog.values = (subject.id, "Slides", True, date(2026, 12, 1))
    monkeypatch.setattr("organizador.controller.BulkFilingDialog", _StubBulkDialog)
    controller.prompt_queue.append(ids[1])

    controller._organise_selection(tuple(ids))

    _pump_until(
        qt_app,
        lambda: "2 organizados" in controller.main_window.inbox_page.import_status_label.text(),
    )
    with controller.database.connect() as connection:
        filed_count = int(
            connection.execute("SELECT COUNT(*) FROM events WHERE action = 'file'").fetchone()[0]
        )
    assert filed_count == 2
    assert controller.database.count_inbox_items() == 1
    pending = controller.database.get_inbox_item(ids[2])
    assert pending is not None
    assert pending.status == "pending"
    assert ids[1] not in controller.prompt_queue
    assert len(controller.database.list_tasks()) == 2
    status = controller.main_window.inbox_page.import_status_label.text()
    assert "2 organizados" in status
    assert "1 com erro" in status

    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def test_filing_runs_off_the_gui_thread(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.inbox_dir / "MAT101_trabalho.pdf"
        path.write_bytes(b"filing content")
        item = controller.database.add_inbox_item(
            path, app_config.downloads_dir / path.name, path.name, path.stat().st_size
        )
        started = threading.Event()
        release = threading.Event()
        real_file_document = controller.filer.file_document

        def slow_file_document(*args: object, **kwargs: object) -> object:
            started.set()
            assert release.wait(10.0)
            return real_file_document(*args, **kwargs)

        monkeypatch.setattr(controller.filer, "file_document", slow_file_document)
        controller._file_item(item.id, subject.id, "Trabalhos", "trabalho.pdf", False, None)

        assert started.wait(5.0)
        assert notices == []
        assert path.exists()
        release.set()
        _pump_until(qt_app, lambda: bool(notices))

        assert notices[0][0][0] == "Ficheiro organizado"
        destination = (
            app_config.university_root / subject.folder_name / "Trabalhos" / "trabalho.pdf"
        )
        assert destination.is_file()
    finally:
        _close_controller(qt_app, controller)


def test_return_ignores_a_second_request_while_in_flight(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database, subject
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.inbox_dir / "devolver.pdf"
        path.write_bytes(b"return content")
        item = controller.database.add_inbox_item(
            path, app_config.downloads_dir / path.name, path.name, path.stat().st_size
        )
        calls = 0
        started = threading.Event()
        release = threading.Event()
        real_return = controller.filer.return_to_origin

        def slow_return(inbox_id: int) -> Path:
            nonlocal calls
            calls += 1
            started.set()
            assert release.wait(10.0)
            return real_return(inbox_id)

        monkeypatch.setattr(controller.filer, "return_to_origin", slow_return)
        controller._return_item(item.id)
        assert started.wait(5.0)
        controller._return_item(item.id)

        release.set()
        _pump_until(qt_app, lambda: any(args[0] == "Ficheiro devolvido" for args, _ in notices))

        assert calls == 1
        assert (app_config.downloads_dir / "devolver.pdf").is_file()
    finally:
        _close_controller(qt_app, controller)


def test_external_file_can_be_organized_and_returned(
    qt_app: QApplication,
    app_config: AppConfig,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    del subject
    origin = tmp_path / "Secretaria"
    origin.mkdir()
    source = origin / "apontamento.pdf"
    source.write_bytes(b"external note content " * 8)
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller.organize_external_paths([str(source)])
        _pump_until(qt_app, lambda: controller.prompt.current_item_id is not None)

        inbox_id = controller.prompt.current_item_id
        assert inbox_id is not None
        item = controller.database.get_inbox_item(inbox_id)
        assert item is not None
        assert item.path.parent == app_config.inbox_dir
        assert item.original_path == source
        assert not source.exists()

        controller.prompt.timer.stop()
        controller.prompt.hide()
        controller._return_item(inbox_id)
        _pump_until(qt_app, lambda: any(args[0] == "Ficheiro devolvido" for args, _ in notices))

        assert source.read_bytes() == b"external note content " * 8
    finally:
        controller.prompt.timer.stop()
        controller.prompt.hide()
        _close_controller(qt_app, controller)


def test_external_organize_refuses_managed_and_unaccepted_files(
    qt_app: QApplication,
    app_config: AppConfig,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        managed = app_config.university_root / subject.folder_name / "Manual.pdf"
        managed.parent.mkdir(parents=True, exist_ok=True)
        managed.write_bytes(b"already filed material " * 8)
        alien = tmp_path / "arquivo.zip"
        alien.write_bytes(b"not a study document " * 8)

        controller.organize_external_paths([str(managed), str(alien)])

        assert controller._pending_transfers == 0
        assert controller.prompt.current_item_id is None
        assert managed.is_file()
        assert alien.is_file()
        assert any(args[0] == "Não foi possível recolher o ficheiro" for args, _ in notices)
    finally:
        _close_controller(qt_app, controller)


def test_undo_ignores_a_second_request_while_in_flight(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.inbox_dir / "MAT101_desfazer.pdf"
        path.write_bytes(b"undo content")
        item = controller.database.add_inbox_item(
            path, app_config.downloads_dir / path.name, path.name, path.stat().st_size
        )
        controller.filer.file_document(item.id, subject.id, "Slides", "desfazer.pdf")
        calls = 0
        started = threading.Event()
        release = threading.Event()
        real_undo = controller.filer.undo_latest_filing

        def slow_undo() -> object:
            nonlocal calls
            calls += 1
            started.set()
            assert release.wait(10.0)
            return real_undo()

        monkeypatch.setattr(controller.filer, "undo_latest_filing", slow_undo)
        controller._undo()
        assert started.wait(5.0)
        controller._undo()

        release.set()
        _pump_until(qt_app, lambda: any(args[0] == "Organização desfeita" for args, _ in notices))

        assert calls == 1
    finally:
        _close_controller(qt_app, controller)


def test_shutdown_transfers_joins_running_workers(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        entered = threading.Event()
        release = threading.Event()

        def slow() -> None:
            entered.set()
            assert release.wait(10.0)
            return None

        controller._submit_transfer("undo", _UndoJob(), slow)
        assert entered.wait(5.0)

        controller._shutdown_transfers(timeout=0.05)

        release.set()
        controller._shutdown_transfers(timeout=5.0)
        with controller._transfer_lock:
            assert not [thread for thread in controller._transfer_threads if thread.is_alive()]
    finally:
        _close_controller(qt_app, controller)


def test_subjects_page_lists_archived_subjects_and_offers_restore(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    database.archive_subject(subject.id)
    window = MainWindow(database, app_config)
    page = window.subjects_page

    page.archived_check.setChecked(True)

    buttons = page.findChildren(QPushButton)
    labels = [item.text() for item in page.findChildren(QLabel)]
    restore = next(control for control in buttons if control.text() == "Restaurar")
    requests: list[int] = []
    page.restore_requested.connect(requests.append)
    restore.click()

    assert requests == [subject.id]
    assert "Arquivada" in labels
    assert "Arquivar" not in {control.text() for control in buttons}
    window.allow_close = True
    window.close()


def test_task_dialog_prefills_and_returns_values(
    qt_app: QApplication, database: Database, subject: Subject
) -> None:
    task = database.add_task(
        "Estudar integrais", subject.id, date(2026, 12, 10), reminder_lead_days=3
    )

    dialog = TaskDialog(task, [subject])

    assert dialog.title_edit.text() == "Estudar integrais"
    assert dialog.subject_combo.currentData() == subject.id
    assert dialog.values == (
        "Estudar integrais",
        subject.id,
        date(2026, 12, 10),
        3,
    )


def test_inbox_rows_without_recovery_are_selectable_and_pruned(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    path = app_config.inbox_dir / "selecionavel.pdf"
    path.write_bytes(b"select me")
    database.add_inbox_item(path, app_config.downloads_dir / path.name, path.name, 100)
    missing = app_config.inbox_dir / "fantasma.pdf"
    database.add_inbox_item(missing, app_config.downloads_dir / missing.name, missing.name, 100)
    window = MainWindow(database, app_config)

    window.inbox_page._selected_ids.add(99999)
    window.inbox_page.refresh()
    checkboxes = window.inbox_page.findChildren(QCheckBox)

    assert 99999 not in window.inbox_page._selected_ids
    assert len(checkboxes) == 2
    window.allow_close = True
    window.close()


def test_task_dialog_keeps_general_tasks_unassigned(
    qt_app: QApplication, database: Database, subject: Subject
) -> None:
    del qt_app
    general = database.add_task("Tarefa geral", None, None)
    assigned = database.add_task("Com disciplina", subject.id, None)

    general_dialog = TaskDialog(general, [subject])
    try:
        assert general_dialog.subject_combo.currentData() is None
        assert general_dialog.values[1] is None
    finally:
        general_dialog.deleteLater()

    assigned_dialog = TaskDialog(assigned, [subject])
    try:
        assert assigned_dialog.subject_combo.currentData() == subject.id
        index = assigned_dialog.subject_combo.findData(None)
        assert index >= 0
        assigned_dialog.subject_combo.setCurrentIndex(index)
        assert assigned_dialog.values[1] is None
    finally:
        assigned_dialog.deleteLater()


def test_filing_prompt_resets_replacement_consent_between_items(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    first_path = app_config.inbox_dir / "primeiro.pdf"
    first_path.write_bytes(b"conteudo do primeiro ficheiro")
    first = database.add_inbox_item(
        first_path,
        app_config.downloads_dir / first_path.name,
        first_path.name,
        first_path.stat().st_size,
    )
    second_path = app_config.inbox_dir / "segundo.pdf"
    second_path.write_bytes(b"conteudo do segundo ficheiro")
    second = database.add_inbox_item(
        second_path,
        app_config.downloads_dir / second_path.name,
        second_path.name,
        second_path.stat().st_size,
    )
    folder = app_config.university_root / subject.folder_name / "Outros"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "primeiro.pdf").write_bytes(b"conteudo do primeiro ficheiro")
    candidate = ExistingDownload.capture(folder / "primeiro.pdf")
    assert candidate is not None
    stored = database.adopt_subject_file(candidate, subject.id, "Outros")
    duplicate = database.get_file(stored.id)
    assert duplicate is not None

    prompt = FilingPrompt(timeout_seconds=30)
    try:
        guess = guess_filing(first.original_name, [subject])
        prompt.show_item(first, [subject], guess, duplicate=duplicate)
        prompt._choose_subject(subject.id, prompt.subject_group.button(subject.id))
        prompt.type_buttons["Outros"].click()
        prompt.replace_check.setChecked(True)
        assert prompt.replace_check.isChecked()

        prompt.show_item(second, [subject], guess, duplicate=duplicate)

        assert not prompt.replace_check.isChecked()
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_filing_prompt_shows_the_recorded_origin_folder(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    tmp_path: Path,
) -> None:
    origin = tmp_path / "Area de Trabalho"
    origin.mkdir()
    source = origin / "apontamento.pdf"
    source.write_bytes(b"conteudo externo " * 4)
    candidate = ExistingDownload.capture(source)
    assert candidate is not None
    item = database.add_inbox_item(source, source, source.name, source.stat().st_size)

    prompt = FilingPrompt(timeout_seconds=30)
    try:
        guess = guess_filing(item.original_name, [subject])
        prompt.show_item(item, [subject], guess)

        assert "Area de Trabalho" in prompt.meta_label.text()
    finally:
        prompt.timer.stop()
        prompt.hide()


def test_subject_files_summary_updates_with_the_document_list(
    qt_app: QApplication,
) -> None:
    del qt_app
    subject = Subject(1, "Cálculo I", "MAT101", "#087A74", (), "MAT101 - Cálculo I", True)
    document = FiledDocument(
        id=7,
        subject_id=1,
        kind="Slides",
        original_name="aula.pdf",
        current_path=Path("C:/uni/MAT101 - Cálculo I/Slides/Aula.pdf"),
        original_path=Path("C:/Downloads/aula.pdf"),
        size=2048,
        filed_at=datetime(2026, 9, 15, 12, tzinfo=UTC),
        indexed_at=None,
    )
    dialog = SubjectFilesDialog(subject, [document], Path("C:/uni/MAT101 - Cálculo I"))
    try:
        assert "1 ficheiro" in dialog.summary_label.text()
        assert "Slides 1" in dialog.kinds_label.text()

        dialog.set_documents([])

        assert "0 ficheiros" in dialog.summary_label.text()
        assert dialog.kinds_label.isHidden()
    finally:
        dialog.deleteLater()


def test_filing_prompt_prefills_from_the_name_template(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    path = app_config.inbox_dir / "MAT101_aula5.pdf"
    path.write_bytes(b"x")
    item = database.add_inbox_item(
        path, app_config.downloads_dir / path.name, path.name, path.stat().st_size
    )
    guess = guess_filing(item.original_name, [subject])
    prompt = FilingPrompt(timeout_seconds=30)

    prompt.show_item(item, [subject], guess)
    assert prompt.name_edit.text() == "MAT101_aula5.pdf"

    prompt.show_item(item, [subject], guess, name_template="{codigo} - {nome_original}")
    assert prompt.name_edit.text() == "MAT101 - MAT101_aula5.pdf"

    prompt.timer.stop()
    prompt.hide()


def test_tasks_page_calendar_marks_days_by_deadline_state(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    today = date.today()
    overdue = database.add_task("Atrasada", subject.id, today - timedelta(days=3))
    today_task = database.add_task("Hoje", subject.id, today)
    future = database.add_task("Futura", subject.id, today + timedelta(days=3))
    completed = database.add_task("Feita", subject.id, today + timedelta(days=1))
    database.set_task_completed(completed.id, True)
    window = MainWindow(database, app_config)
    page = window.tasks_page
    open_formats = {
        day: page.calendar.dateTextFormat(QDate(day.year, day.month, day.day))
        for day in (overdue.due_date, today_task.due_date, future.due_date)
    }

    assert open_formats[overdue.due_date].fontWeight() == QFont.Weight.Bold
    assert open_formats[today_task.due_date].fontWeight() == QFont.Weight.Bold
    assert open_formats[future.due_date].fontWeight() == QFont.Weight.Bold
    assert open_formats[overdue.due_date].foreground().color().name() == "#ff818b"
    assert open_formats[today_task.due_date].foreground().color().name() == "#f1bb68"
    assert open_formats[future.due_date].foreground().color().name() == "#49cfc0"
    assert open_formats[overdue.due_date].background().color().name() == "#4a262e"
    assert open_formats[today_task.due_date].background().color().name() == "#4a3520"
    assert open_formats[future.due_date].background().color().name() == "#1e4d46"
    completed_format = page.calendar.dateTextFormat(
        QDate(completed.due_date.year, completed.due_date.month, completed.due_date.day)
    )
    assert completed_format.fontWeight() != QFont.Weight.Bold
    assert completed_format.foreground().color().name() == "#9baabd"
    assert completed_format.background().color().name() == "#222f3e"
    window.allow_close = True
    window.close()


def test_tasks_page_calendar_click_filters_and_toggle_clears(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    today = date.today()
    database.add_task("Um", subject.id, today)
    database.add_task("Dois", subject.id, today)
    database.add_task("Outro dia", subject.id, today + timedelta(days=5))
    window = MainWindow(database, app_config)
    page = window.tasks_page

    page.calendar.clicked.emit(QDate(today.year, today.month, today.day))

    assert page._selected_date == today
    assert not page.clear_filter_button.isHidden()
    page.refresh()
    titles = [label_.text() for label_ in page.findChildren(QLabel)]
    assert "Um" in titles
    assert "Dois" in titles
    assert "Outro dia" not in titles

    page.calendar.clicked.emit(QDate(today.year, today.month, today.day))

    assert page._selected_date is None
    window.allow_close = True
    window.close()


def test_tasks_page_calendar_activation_prefills_the_deadline(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    chosen = date.today().replace(day=20)
    window = MainWindow(database, app_config)
    page = window.tasks_page
    page.due_check.setChecked(False)

    page.calendar.activated.emit(QDate(chosen.year, chosen.month, chosen.day))

    assert page.due_check.isChecked()
    assert page.due_edit.date() == QDate(chosen.year, chosen.month, chosen.day)
    window.allow_close = True
    window.close()


def test_subjects_page_shows_counts_and_opens_files_overview(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    filer: FilingService,
) -> None:
    download = app_config.downloads_dir / "revolucao.pdf"
    download.write_bytes(b"material historico" * 12)
    item = filer.ingest(download)
    assert item is not None
    document = filer.file_document(item.id, subject.id, "Slides", download.name)
    database.add_subject("Sem ficheiros", "", "#123456", (), "SEM - Sem ficheiros")
    window = MainWindow(database, app_config)
    page = window.subjects_page
    row_labels = [text.text() for text in page.findChildren(QLabel)]
    assert any("1 ficheiro ·" in text for text in row_labels)
    assert any("Ainda sem ficheiros organizados" in text for text in row_labels)

    requests: list[int] = []
    page.view_files_requested.connect(requests.append)
    view_button = next(
        control for control in page.findChildren(QPushButton) if control.text() == "Ver ficheiros"
    )
    view_button.click()
    assert requests == [subject.id]

    folder_path = app_config.university_root / subject.folder_name
    dialog = SubjectFilesDialog(subject, [document], folder_path)
    opened: list[object] = []
    dialog.open_requested.connect(opened.append)
    dialog_labels = [text.text() for text in dialog.findChildren(QLabel)]
    assert document.current_path.name in dialog_labels
    assert any("1 ficheiro ·" in text for text in dialog_labels)
    assert any("Slides 1" in text for text in dialog_labels)
    dialog_buttons = dialog.findChildren(QPushButton)
    next(control for control in dialog_buttons if control.text() == "Abrir").click()
    next(control for control in dialog_buttons if control.text() == "Abrir pasta").click()
    assert opened == [document.current_path, folder_path]

    empty_dialog = SubjectFilesDialog(subject, [], folder_path)
    empty_labels = [text.text() for text in empty_dialog.findChildren(QLabel)]
    assert any("Ainda não há ficheiros organizados" in text for text in empty_labels)
    window.allow_close = True
    window.close()


def test_controller_view_subject_files_opens_dialog_without_touching_files(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    del database, qt_app
    controller = AppController(app_config)
    shown: list[list[object]] = []

    def capture_exec(dialog: object) -> QDialog.DialogCode:
        documents = getattr(dialog, "documents", ())
        shown.append(list(documents))
        return QDialog.DialogCode.Rejected

    monkeypatch.setattr(SubjectFilesDialog, "exec", capture_exec)
    controller._view_subject_files(subject.id)

    assert shown == [[]]
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()


def _update_info(version: tuple[int, int, int] = (0, 6, 2)) -> updater.UpdateInfo:
    return updater.UpdateInfo(
        version=version,
        tag_name=f"v{version[0]}.{version[1]}.{version[2]}",
        zip_url="https://example.invalid/organizador.zip",
        sha256_url="https://example.invalid/organizador.zip.sha256",
    )


def _watched_controller(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[AppController, list[tuple[tuple[object, ...], dict[str, object]]]]:
    controller = AppController(app_config)
    notices: list[tuple[tuple[object, ...], dict[str, object]]] = []
    monkeypatch.setattr(
        controller.tray,
        "notify",
        lambda *args, **kwargs: notices.append((args, kwargs)),
    )
    monkeypatch.setattr(TrayIcon, "available", property(lambda self: True))
    qt_app.processEvents()
    return controller, notices


def _close_controller(qt_app: QApplication, controller: AppController) -> None:
    controller._shutdown_transfers()
    controller.indexer.shutdown()
    controller.tray.hide()
    controller.main_window.allow_close = True
    controller.main_window.close()
    qt_app.processEvents()


def _pump_until(qt_app: QApplication, predicate: Callable[[], bool], timeout: float = 10.0) -> None:
    """Pump Qt events until a background transfer delivers its outcome."""

    deadline = time.monotonic() + timeout
    while not predicate():
        if time.monotonic() >= deadline:
            raise AssertionError("timed out waiting for a background file transfer")
        qt_app.processEvents()
        time.sleep(0.01)


def _finish_check(
    controller: AppController, result: updater.UpdateCheckResult, automatic: bool
) -> None:
    controller._on_update_check_finished(result, automatic, controller._update_check_generation)


def test_intake_notice_single_file_keeps_named_message(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._queue_intake_notice("apontamentos.pdf")
        controller._flush_intake_notices()

        assert notices == [
            (
                (
                    "Novo material na Caixa de Entrada",
                    "apontamentos.pdf está pronto para organizar.",
                ),
                {},
            )
        ]
    finally:
        _close_controller(qt_app, controller)


def test_intake_notices_batch_into_one_summary(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._queue_intake_notice("a.pdf")
        controller._queue_intake_notice("b.pdf")
        controller._queue_intake_notice("c.pdf")
        controller._flush_intake_notices()

        assert notices == [
            (
                ("Novo material na Caixa de Entrada", "3 ficheiros estão prontos para organizar."),
                {},
            )
        ]
    finally:
        _close_controller(qt_app, controller)


def test_first_filed_notice_is_immediate(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._notify_filed("apontamentos.pdf", "Matemática / Exercícios")

        assert notices == [
            (
                (
                    "Ficheiro organizado",
                    "apontamentos.pdf foi guardado em Matemática / Exercícios.",
                ),
                {},
            )
        ]
    finally:
        _close_controller(qt_app, controller)


def test_filed_notices_aggregate_into_one_summary(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._last_filed_notice = time.monotonic()
        controller._notify_filed("a.pdf", "Matemática / Exercícios")
        controller._notify_filed("b.pdf", "Matemática / Exercícios")
        controller._notify_filed("c.pdf", "Biologia / Apontamentos")
        controller._flush_filed_notices()

        assert notices == [(("Ficheiros organizados", "3 ficheiros organizados"), {})]
    finally:
        _close_controller(qt_app, controller)


def test_quiet_intake_suppresses_intake_and_filed_notices(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller.config.quiet_intake = True
        controller._queue_intake_notice("a.pdf")
        controller._flush_intake_notices()
        controller._notify_filed("a.pdf", "Matemática / Exercícios")
        controller._flush_filed_notices()

        assert notices == []
    finally:
        _close_controller(qt_app, controller)


def test_begin_update_check_skips_while_installing(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._update_installing = True

        controller._begin_update_check(automatic=False)

        assert controller._update_checking is False
        assert notices == []
    finally:
        _close_controller(qt_app, controller)


def test_begin_update_check_reports_dev_mode_without_freezing(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._begin_update_check(automatic=False)

        assert controller._update_checking is False
        assert notices and notices[0][0][0] == "Sem atualizações nesta instalação"
    finally:
        _close_controller(qt_app, controller)


def test_update_check_no_update_clears_pending_and_reports(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._pending_update = _update_info()

        _finish_check(
            controller,
            updater.UpdateCheckResult(updater.UpdateCheckStatus.NO_UPDATE),
            False,
        )

        assert controller._pending_update is None
        assert controller._update_checking is False
        assert not controller.tray.install_update_action.isVisible()
        assert notices and notices[0][0][0] == "Sem atualizações"
    finally:
        _close_controller(qt_app, controller)


def test_update_check_error_keeps_pending_and_warns_only_manual(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        pending = _update_info()
        controller._pending_update = pending

        _finish_check(
            controller,
            updater.UpdateCheckResult(updater.UpdateCheckStatus.ERROR, error="boom"),
            False,
        )

        assert controller._pending_update is pending
        assert controller.tray.install_update_action.isVisible()
        assert notices and notices[0][0][0] == "Não foi possível procurar atualizações."
        assert notices[0][1].get("icon") == QSystemTrayIcon.MessageIcon.Warning

        _finish_check(
            controller,
            updater.UpdateCheckResult(updater.UpdateCheckStatus.ERROR, error="boom"),
            True,
        )

        assert len(notices) == 1
    finally:
        _close_controller(qt_app, controller)


def test_update_check_available_sets_pending_install_action(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        info = _update_info()

        _finish_check(
            controller,
            updater.UpdateCheckResult(updater.UpdateCheckStatus.UPDATE_AVAILABLE, update=info),
            True,
        )

        assert controller._pending_update is info
        assert controller.tray.install_update_action.isVisible()
        assert "0.6.2" in controller.tray.install_update_action.text()
        assert notices and notices[0][0][0] == "Atualização disponível"
    finally:
        _close_controller(qt_app, controller)


def test_update_install_finished_failure_restores_install_action(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._pending_update = _update_info()
        controller._update_installing = True

        controller._on_update_install_finished("boom")

        assert controller._update_installing is False
        assert controller.tray.install_update_action.isVisible()
        assert "0.6.2" in controller.tray.install_update_action.text()
        assert notices and notices[0][0][0] == "Atualização falhou"
    finally:
        _close_controller(qt_app, controller)


def test_update_install_finished_abort_stays_quiet(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        controller._update_installing = True

        controller._on_update_install_finished(None)

        assert controller._update_installing is False
        assert notices == []
    finally:
        _close_controller(qt_app, controller)


def test_update_install_finished_transaction_launches_helper_before_shutdown(
    qt_app: QApplication,
    app_config: AppConfig,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    app = tmp_path / "Prog App"
    (app / "_internal").mkdir(parents=True)
    (app / "Organizador.exe").write_bytes(b"old")
    transaction = updater.create_update_transaction(app, "0.6.2", data_dir=app_config.data_dir)
    updater.write_update_helper(transaction)
    launched: dict[str, object] = {}
    shutdowns: list[bool] = []

    class _Process:
        pass

    def fake_launch_helper(
        launched_transaction: updater.UpdateTransaction, *, wait_ready: bool = True
    ) -> _Process:
        launched["transaction"] = launched_transaction
        launched["wait_ready"] = wait_ready
        return _Process()

    monkeypatch.setattr(updater, "launch_update_helper", fake_launch_helper)
    monkeypatch.setattr(updater, "helper_ready_received", lambda _transaction: True)
    monkeypatch.setattr(controller, "shutdown", lambda: shutdowns.append(True))
    try:
        controller._update_installing = True

        controller._on_update_install_finished(transaction)
        qt_app.processEvents()

        assert launched["transaction"] is transaction
        assert launched["wait_ready"] is False
        assert controller._update_restart_armed is True
        assert shutdowns == [True]
    finally:
        updater.abort_update_transaction(transaction)
        _close_controller(qt_app, controller)


def test_legacy_rollback_bridge_retains_then_cleans(
    qt_app: QApplication,
    app_config: AppConfig,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    app = tmp_path / "Legacy App"
    (app / "_internal").mkdir(parents=True)
    (app / "Organizador.exe").write_bytes(b"current")
    old = tmp_path / "Organizador.old"
    (old / "_internal").mkdir(parents=True)
    (old / "Organizador.exe").write_bytes(b"previous")
    monkeypatch.setattr(updater, "app_directory", lambda: app)
    try:
        controller._handle_legacy_rollback_bridge()

        assert old.is_dir()
        marker = app_config.data_dir / "updates" / "legacy-rollback-retained"
        assert marker.is_file()

        controller._handle_legacy_rollback_bridge()
        deadline = time.monotonic() + 5.0
        while old.exists() and time.monotonic() < deadline:
            qt_app.processEvents()
            time.sleep(0.025)

        assert not old.exists()
    finally:
        _close_controller(qt_app, controller)


def _write_result(
    app_config: AppConfig,
    status: updater.UpdateResultStatus,
    transaction_id: str = "a" * 32,
) -> None:
    state_dir = app_config.data_dir / "updates" / transaction_id
    state_dir.mkdir(parents=True)
    updater.write_update_result(
        state_dir / "result.json",
        updater.UpdateResult(
            transaction_id=transaction_id,
            status=status,
            phase="complete",
            committed=status is not updater.UpdateResultStatus.ROLLED_BACK,
            rollback_succeeded=True,
            error=None,
            old_pid=1,
            new_pid=2,
            started_at="2026-09-03T00:00:00+00:00",
            finished_at="2026-09-03T00:00:01+00:00",
            app_dir=state_dir,
            rollback_dir=state_dir,
        ),
    )


def test_pending_update_success_result_is_shown_once(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        _write_result(app_config, updater.UpdateResultStatus.SUCCEEDED)

        controller._show_pending_update_result()
        controller._show_pending_update_result()

        assert len(notices) == 1
        assert notices[0][0][0] == "Atualização instalada"
        assert "com sucesso" in notices[0][0][1]
    finally:
        _close_controller(qt_app, controller)


def test_pending_update_rollback_result_warns_once(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        _write_result(app_config, updater.UpdateResultStatus.ROLLED_BACK)

        controller._show_pending_update_result()
        controller._show_pending_update_result()

        assert len(notices) == 1
        assert notices[0][0][0] == "Atualização revertida"
        assert notices[0][1].get("icon") == QSystemTrayIcon.MessageIcon.Warning
    finally:
        _close_controller(qt_app, controller)


def test_subject_files_dialog_offers_reindex_for_failed_documents(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    del qt_app
    path = app_config.university_root / subject.folder_name / "Outros" / "quebrado.docx"
    path.write_bytes(b"not an OOXML archive")
    candidate = ExistingDownload.capture(path)
    assert candidate is not None
    filed = database.adopt_subject_file(candidate, subject.id, "Outros")
    indexer_documents = database.get_file(filed.id)
    assert indexer_documents is not None
    from organizador.indexer import DocumentIndexer

    indexer = DocumentIndexer(database)
    indexer.index_document(indexer_documents)
    indexer.shutdown()
    stored = database.get_file(filed.id)
    assert stored is not None
    assert stored.index_state == "failed"

    dialog = SubjectFilesDialog(subject, [dataclasses.replace(stored)], app_config.university_root)
    requested: list[int] = []
    dialog.reindex_requested.connect(requested.append)
    buttons = {control.text(): control for control in dialog.findChildren(QPushButton)}
    labels = {control.text() for control in dialog.findChildren(QLabel)}

    assert "Reindexar" in buttons
    assert any("pesquisável pelo nome" in text for text in labels)
    buttons["Reindexar"].click()

    assert requested == [filed.id]
    dialog.allow_close = True
    dialog.close()


def test_subject_files_dialog_offers_reindex_for_healthy_documents(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
    subject: Subject,
) -> None:
    del qt_app
    path = app_config.university_root / subject.folder_name / "Outros" / "saudavel.txt"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("derivadas saudáveis", encoding="utf-8")
    candidate = ExistingDownload.capture(path)
    assert candidate is not None
    filed = database.adopt_subject_file(candidate, subject.id, "Outros")
    document = database.get_file(filed.id)
    assert document is not None
    database.replace_document_pages(
        filed.id, subject.name, path.name, ["derivadas saudáveis"], expected_path=path
    )
    stored = database.get_file(filed.id)
    assert stored is not None
    assert stored.index_state == ""

    dialog = SubjectFilesDialog(subject, [dataclasses.replace(stored)], app_config.university_root)
    requested: list[int] = []
    dialog.reindex_requested.connect(requested.append)
    buttons = {control.text(): control for control in dialog.findChildren(QPushButton)}

    assert "Reindexar" in buttons
    buttons["Reindexar"].click()

    assert requested == [filed.id]
    dialog.allow_close = True
    dialog.close()


def _search_combo_texts(combo: QComboBox) -> list[str]:
    return [combo.itemText(index) for index in range(combo.count())]


def test_search_filter_combos_populate_and_filter_results(
    qt_app: QApplication,
    app_config: AppConfig,
    database: Database,
) -> None:
    first = database.add_subject("Cálculo I", "MAT101", "#000000", (), "MAT101")
    second = database.add_subject("Física I", "FIS101", "#000000", (), "FIS101")

    def file_record(name: str, subject_id: int, kind: str, text: str) -> None:
        inbox_path = app_config.inbox_dir / name
        inbox_path.write_text(text, encoding="utf-8")
        item = database.add_inbox_item(
            inbox_path, app_config.downloads_dir / name, name, inbox_path.stat().st_size
        )
        destination = app_config.university_root / name
        inbox_path.replace(destination)
        filed = database.record_filing(item.id, subject_id, kind, destination)
        database.replace_document_pages(filed.id, name, name, (text,), expected_path=destination)

    file_record("derivadas.txt", first.id, "Slides", "derivadas")
    file_record("integral.txt", first.id, "Testes", "derivadas")
    file_record("cinematica.txt", second.id, "Slides", "derivadas")
    window = MainWindow(database, app_config)
    try:
        page = window.search_page
        assert _search_combo_texts(page.subject_combo) == [
            "Todas as disciplinas",
            "Cálculo I (MAT101)",
            "Física I (FIS101)",
        ]
        assert _search_combo_texts(page.kind_combo) == [
            "Todos os tipos",
            "Slides",
            "Exercícios",
            "Testes",
            "Trabalhos",
            "Outros",
        ]

        page.search_edit.setText("derivadas")
        page.kind_combo.setCurrentIndex(page.kind_combo.findData("Testes"))
        qt_app.processEvents()

        assert "1 resultado" in page.status_label.text()
        titles = [
            row.findChild(QLabel).text()
            for row in page.results_layout.parentWidget().findChildren(QFrame)
            if row.objectName() == "ListRow"
        ]
        assert titles == ["integral.txt"]

        page.search_edit.clear()
        page.kind_combo.setCurrentIndex(0)
        page.subject_combo.setCurrentIndex(page.subject_combo.findData(second.id))
        qt_app.processEvents()

        assert page.status_label.text() == "1 documento"
        browse_titles = [
            row.findChild(QLabel).text()
            for row in page.results_layout.parentWidget().findChildren(QFrame)
            if row.objectName() == "ListRow"
        ]
        assert browse_titles == ["cinematica.txt"]
    finally:
        window.allow_close = True
        window.close()


def test_retry_failed_indexes_requeues_documents_through_the_worker(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from docx import Document

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.university_root / "recuperar.docx"
        path.write_bytes(b"not an OOXML archive")
        candidate = ExistingDownload.capture(path)
        assert candidate is not None
        subject = controller.database.add_subject(
            "Recuperação", "REC", "#123456", (), "Recuperação"
        )
        filed = controller.database.adopt_subject_file(candidate, subject.id, "Outros")
        controller.indexer.index_document(controller.database.get_file(filed.id))
        failed = controller.database.get_file(filed.id)
        assert failed is not None
        assert failed.index_state == "failed"

        repaired = Document()
        repaired.add_paragraph("conteúdo recuperado pelo worker")
        repaired.save(path)
        controller._retry_failed_indexes()

        deadline = time.monotonic() + 15.0
        while not controller.database.search("recuperado") and time.monotonic() < deadline:
            qt_app.processEvents()
            time.sleep(0.05)

        assert controller.database.search("recuperado")
        assert controller.database.list_failed_index_documents() == []
    finally:
        _close_controller(qt_app, controller)


def test_reindexar_marks_the_row_and_refreshes_after_indexing(
    qt_app: QApplication,
    app_config: AppConfig,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.university_root / subject.folder_name / "Outros" / "reindexar.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("conteudo reindexado de teste", encoding="utf-8")
        candidate = ExistingDownload.capture(path)
        assert candidate is not None
        filed = controller.database.adopt_subject_file(candidate, subject.id, "Outros")
        document = controller.database.get_file(filed.id)
        assert document is not None
        dialog = SubjectFilesDialog(subject, [document], app_config.university_root)
        try:
            controller._subject_files_dialog = dialog
            controller._reindex_document(document.id)

            assert not dialog.reindex_buttons[document.id].isEnabled()
            assert dialog.pending_labels[document.id].text() == "A reindexar…"
            assert not dialog.pending_labels[document.id].isHidden()

            def refined() -> bool:
                note = dialog.pending_labels.get(document.id)
                button = dialog.reindex_buttons.get(document.id)
                return (
                    note is not None
                    and note.isHidden()
                    and button is not None
                    and button.isEnabled()
                )

            _pump_until(qt_app, refined)

            assert controller.database.search("reindexado")
        finally:
            controller._subject_files_dialog = None
            dialog.deleteLater()
            qt_app.processEvents()
    finally:
        _close_controller(qt_app, controller)


def test_reindexar_reports_a_busy_queue(
    qt_app: QApplication,
    app_config: AppConfig,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.university_root / subject.folder_name / "Outros" / "ocupada.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("conteudo na fila", encoding="utf-8")
        candidate = ExistingDownload.capture(path)
        assert candidate is not None
        filed = controller.database.adopt_subject_file(candidate, subject.id, "Outros")
        document = controller.database.get_file(filed.id)
        assert document is not None
        dialog = SubjectFilesDialog(subject, [document], app_config.university_root)
        try:
            controller._subject_files_dialog = dialog
            monkeypatch.setattr(controller.indexer, "reindex", lambda _document: False)

            controller._reindex_document(document.id)

            assert "fila de indexação" in dialog.pending_labels[document.id].text()
        finally:
            controller._subject_files_dialog = None
            dialog.deleteLater()
            qt_app.processEvents()
    finally:
        _close_controller(qt_app, controller)


def test_subject_files_dialog_offers_move(qt_app: QApplication) -> None:
    del qt_app
    subject = Subject(1, "Cálculo I", "MAT101", "#087A74", (), "MAT101 - Cálculo I", True)
    document = FiledDocument(
        id=7,
        subject_id=1,
        kind="Slides",
        original_name="aula.pdf",
        current_path=Path("C:/uni/MAT101 - Cálculo I/Slides/Aula.pdf"),
        original_path=Path("C:/Downloads/aula.pdf"),
        size=200,
        filed_at=datetime(2026, 9, 15, 12, tzinfo=UTC),
        indexed_at=None,
    )
    dialog = SubjectFilesDialog(subject, [document], Path("C:/uni/MAT101 - Cálculo I"))
    try:
        requested: list[int] = []
        dialog.move_requested.connect(requested.append)
        buttons = {control.text(): control for control in dialog.findChildren(QPushButton)}

        assert "Mover" in buttons
        buttons["Mover"].click()

        assert requested == [document.id]
    finally:
        dialog.deleteLater()


def test_move_dialog_requires_a_real_destination(qt_app: QApplication) -> None:
    del qt_app
    subject = Subject(1, "Cálculo I", "MAT101", "#087A74", (), "MAT101 - Cálculo I", True)
    other = Subject(2, "Física", "FIS110", "#3C64A3", (), "FIS110 - Física", True)
    document = FiledDocument(
        id=7,
        subject_id=1,
        kind="Slides",
        original_name="aula.pdf",
        current_path=Path("C:/uni/MAT101 - Cálculo I/Slides/Aula.pdf"),
        original_path=Path("C:/Downloads/aula.pdf"),
        size=200,
        filed_at=datetime(2026, 9, 15, 12, tzinfo=UTC),
        indexed_at=None,
    )
    dialog = MoveDocumentDialog(document, [subject, other])
    try:
        assert not dialog.move_button.isEnabled()
        assert "outro tipo" in dialog.hint_label.text()

        other_kind = dialog.kind_combo.findText("Trabalhos")
        dialog.kind_combo.setCurrentIndex(other_kind)
        assert dialog.move_button.isEnabled()
        assert dialog.hint_label.isHidden()

        dialog.subject_combo.setCurrentIndex(dialog.subject_combo.findData(other.id))
        dialog.kind_combo.setCurrentIndex(dialog.kind_combo.findText("Slides"))

        assert dialog.selection() == (other.id, "Slides")
    finally:
        dialog.deleteLater()


def test_controller_moves_a_document_through_the_worker(
    qt_app: QApplication,
    app_config: AppConfig,
    subject: Subject,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        other = controller.database.add_subject(
            "Física Geral", "FIS110", "#3C64A3", (), "FIS110 - Física Geral"
        )
        controller.filer.ensure_subject_structure(other)
        path = app_config.university_root / subject.folder_name / "Outros" / "mover.txt"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text("conteudo para mover", encoding="utf-8")
        candidate = ExistingDownload.capture(path)
        assert candidate is not None
        filed = controller.database.adopt_subject_file(candidate, subject.id, "Outros")

        monkeypatch.setattr(MoveDocumentDialog, "exec", lambda self: QDialog.DialogCode.Accepted)
        monkeypatch.setattr(MoveDocumentDialog, "selection", lambda self: (other.id, "Trabalhos"))

        controller._open_move_dialog(filed.id)
        target = app_config.university_root / other.folder_name / "Trabalhos" / "mover.txt"

        _pump_until(qt_app, lambda: any(args[0] == "Ficheiro movido" for args, _ in notices))

        assert target.is_file()
        assert not path.exists()
        stored = controller.database.get_file(filed.id)
        assert stored is not None
        assert stored.subject_id == other.id
        assert stored.kind == "Trabalhos"
    finally:
        _close_controller(qt_app, controller)


def test_restore_relaunch_keeps_a_custom_data_directory(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import subprocess

    from organizador import updater

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        commands: list[list[str]] = []
        monkeypatch.setattr(updater, "is_frozen", lambda: True)
        monkeypatch.setattr(subprocess, "Popen", lambda args, **kwargs: commands.append(list(args)))
        monkeypatch.setattr(controller, "shutdown", lambda: None)

        controller._relaunch_for_restore()

        assert len(commands) == 1
        command = commands[0][-1]
        assert "--data-dir" in command
        assert str(app_config.data_dir.resolve()) in command
        assert f'"{app_config.data_dir.resolve()}"' in command

        commands.clear()
        monkeypatch.setattr("organizador.controller.default_data_dir", lambda: app_config.data_dir)
        controller._relaunch_for_restore()

        assert len(commands) == 1
        assert "--data-dir" not in commands[0][-1]
    finally:
        _close_controller(qt_app, controller)


def test_stale_check_result_is_ignored_and_install_waits_for_quiet(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        info = _update_info()
        controller._pending_update = info
        controller._update_checking = True

        controller._install_pending_update()

        assert controller._update_installing is False
        assert notices == []

        controller._update_check_generation = 5
        controller._on_update_check_finished(
            updater.UpdateCheckResult(updater.UpdateCheckStatus.NO_UPDATE), False, 3
        )

        assert controller._pending_update is info
        assert controller._update_checking is True
    finally:
        _close_controller(qt_app, controller)


def test_downloads_during_update_preparation_are_deferred_and_released(
    qt_app: QApplication,
    app_config: AppConfig,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        path = app_config.downloads_dir / "durante-atualizacao.pdf"
        path.write_bytes(b"download during update preparation")
        controller._update_installing = True

        controller._ingest_download(controller._watcher_generation, path)
        controller._ingest_download(controller._watcher_generation, path)

        assert controller.database.count_inbox_items() == 0
        assert path.is_file()
        assert [candidate for _, candidate in controller._deferred_downloads] == [path]

        controller._on_update_install_finished("falha simulada")
        _pump_until(qt_app, lambda: not path.exists())

        assert controller.database.count_inbox_items() == 1
        assert controller._deferred_downloads == []
    finally:
        _close_controller(qt_app, controller)


def test_handshake_activation_failure_restores_pending_migration(
    qt_app: QApplication,
    app_config: AppConfig,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from organizador.controller import StartupState
    from organizador.recovery import RecoveryCoordinator

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        with controller.database.connect() as connection:
            connection.execute("ALTER TABLE tasks DROP COLUMN reminder_lead_days")
            connection.commit()
        coordinator = RecoveryCoordinator(app_config.data_dir)
        bundle = coordinator.prepare_migration()
        assert bundle is not None

        app = tmp_path / "Handshake App"
        (app / "_internal").mkdir(parents=True)
        (app / "Organizador.exe").write_bytes(b"candidate")
        transaction = updater.create_update_transaction(app, "0.6.3", data_dir=app_config.data_dir)
        state = StartupState(configured=True, services_ready=True)
        exits: list[int] = []
        monkeypatch.setattr(QApplication, "exit", lambda code=0: exits.append(code))
        monkeypatch.setattr(
            QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.StandardButton.Ok
        )

        def fail_activate(
            _state: StartupState, *, background: bool = False, smoke_test: bool = False
        ) -> None:
            raise RuntimeError("activation exploded")

        monkeypatch.setattr(controller, "activate", fail_activate)

        controller._commit_update_handshake(
            transaction, bundle, coordinator, state, background=True
        )

        assert exits == [1]
        assert not (bundle.path / "healthy").exists()
        assert coordinator.restore_pending() is None
        inspection = controller.database.inspect_schema()
        assert "tasks.reminder_lead_days" in inspection.missing_additions
    finally:
        updater.abort_update_transaction(transaction)
        _close_controller(qt_app, controller)


def test_handshake_acknowledges_health_before_closing_data_rollback(
    qt_app: QApplication,
    app_config: AppConfig,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from organizador.controller import StartupState
    from organizador.recovery import RecoveryCoordinator

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        with controller.database.connect() as connection:
            connection.execute("ALTER TABLE tasks DROP COLUMN reminder_lead_days")
            connection.commit()
        coordinator = RecoveryCoordinator(app_config.data_dir)
        bundle = coordinator.prepare_migration()
        assert bundle is not None
        controller.database.initialize()

        app = tmp_path / "Handshake App"
        (app / "_internal").mkdir(parents=True)
        (app / "Organizador.exe").write_bytes(b"candidate")
        transaction = updater.create_update_transaction(app, "0.6.3", data_dir=app_config.data_dir)
        state = StartupState(configured=True, services_ready=True)
        order: list[str] = []

        def noop_activate(
            _state: StartupState, *, background: bool = False, smoke_test: bool = False
        ) -> None:
            return None

        real_validate = coordinator.validate_migrated

        def record_validate(bundle_arg: object) -> None:
            order.append("validate")
            real_validate(bundle_arg)  # type: ignore[arg-type]

        monkeypatch.setattr(controller, "activate", noop_activate)
        monkeypatch.setattr(coordinator, "validate_migrated", record_validate)
        monkeypatch.setattr(
            updater, "mark_update_healthy", lambda *args, **kwargs: order.append("ack")
        )
        monkeypatch.setattr(coordinator, "mark_healthy", lambda *args, **kwargs: order.append("db"))
        monkeypatch.setattr(QApplication, "exit", lambda code=0: order.append("exit"))
        monkeypatch.setattr(
            QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.StandardButton.Ok
        )

        controller._commit_update_handshake(
            transaction, bundle, coordinator, state, background=True
        )

        assert order == ["validate", "ack", "db"]
    finally:
        updater.abort_update_transaction(transaction)
        _close_controller(qt_app, controller)


def test_handshake_validation_failure_skips_ack_and_restores(
    qt_app: QApplication,
    app_config: AppConfig,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from organizador.controller import StartupState
    from organizador.recovery import RecoveryCoordinator

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        with controller.database.connect() as connection:
            connection.execute("ALTER TABLE tasks DROP COLUMN reminder_lead_days")
            connection.commit()
        coordinator = RecoveryCoordinator(app_config.data_dir)
        bundle = coordinator.prepare_migration()
        assert bundle is not None

        app = tmp_path / "Handshake App"
        (app / "_internal").mkdir(parents=True)
        (app / "Organizador.exe").write_bytes(b"candidate")
        transaction = updater.create_update_transaction(app, "0.6.3", data_dir=app_config.data_dir)
        state = StartupState(configured=True, services_ready=True)
        exits: list[int] = []
        acknowledged: list[bool] = []
        marked: list[bool] = []
        restored: list[bool] = []

        def noop_activate(
            _state: StartupState, *, background: bool = False, smoke_test: bool = False
        ) -> None:
            return None

        def fail_validation(*args: object, **kwargs: object) -> None:
            raise RuntimeError("migração inválida")

        monkeypatch.setattr(controller, "activate", noop_activate)
        monkeypatch.setattr(coordinator, "validate_migrated", fail_validation)
        monkeypatch.setattr(
            updater, "mark_update_healthy", lambda *args, **kwargs: acknowledged.append(True)
        )
        monkeypatch.setattr(
            coordinator, "mark_healthy", lambda *args, **kwargs: marked.append(True)
        )
        monkeypatch.setattr(coordinator, "restore_pending", lambda: restored.append(True) or None)
        monkeypatch.setattr(QApplication, "exit", lambda code=0: exits.append(code))
        monkeypatch.setattr(
            QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.StandardButton.Ok
        )

        controller._commit_update_handshake(
            transaction, bundle, coordinator, state, background=True
        )

        assert exits == [1]
        assert restored == [True]
        assert acknowledged == []
        assert marked == []
    finally:
        updater.abort_update_transaction(transaction)
        _close_controller(qt_app, controller)


def test_handshake_ack_failure_restores_data_rollback(
    qt_app: QApplication,
    app_config: AppConfig,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from organizador.controller import StartupState
    from organizador.recovery import RecoveryCoordinator

    controller, _notices = _watched_controller(qt_app, app_config, monkeypatch)
    try:
        with controller.database.connect() as connection:
            connection.execute("ALTER TABLE tasks DROP COLUMN reminder_lead_days")
            connection.commit()
        coordinator = RecoveryCoordinator(app_config.data_dir)
        bundle = coordinator.prepare_migration()
        assert bundle is not None

        app = tmp_path / "Handshake App"
        (app / "_internal").mkdir(parents=True)
        (app / "Organizador.exe").write_bytes(b"candidate")
        transaction = updater.create_update_transaction(app, "0.6.3", data_dir=app_config.data_dir)
        state = StartupState(configured=True, services_ready=True)
        exits: list[int] = []
        marked: list[bool] = []
        restored: list[bool] = []

        def noop_activate(
            _state: StartupState, *, background: bool = False, smoke_test: bool = False
        ) -> None:
            return None

        def fail_ack(*args: object, **kwargs: object) -> None:
            raise OSError("disco cheio")

        monkeypatch.setattr(controller, "activate", noop_activate)
        monkeypatch.setattr(coordinator, "validate_migrated", lambda *args, **kwargs: None)
        monkeypatch.setattr(updater, "mark_update_healthy", fail_ack)
        monkeypatch.setattr(
            coordinator, "mark_healthy", lambda *args, **kwargs: marked.append(True)
        )
        monkeypatch.setattr(coordinator, "restore_pending", lambda: restored.append(True) or None)
        monkeypatch.setattr(QApplication, "exit", lambda code=0: exits.append(code))
        monkeypatch.setattr(
            QMessageBox, "critical", lambda *args, **kwargs: QMessageBox.StandardButton.Ok
        )

        controller._commit_update_handshake(
            transaction, bundle, coordinator, state, background=True
        )

        assert exits == [1]
        assert restored == [True]
        assert marked == []
    finally:
        updater.abort_update_transaction(transaction)
        _close_controller(qt_app, controller)
